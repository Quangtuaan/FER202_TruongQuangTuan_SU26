from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "Week10" / "Week10_BaoCaoNhom_HoanChinh.docx"
NAVY = "203748"
BLUE = "2E74B5"
GOLD = "C5A880"
LIGHT = "E8EEF5"
PALE = "F4F6F9"
GRAY = "5B6570"
WHITE = "FFFFFF"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc.get_or_add_tcPr()
    tc_mar = tc.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc.append(tc_mar)
    for tag, val in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")


def set_cell_text(cell, text, bold=False, color=None, size=9.5, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    r = p.add_run(str(text))
    font(r, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    margins(cell)


def font(run, name="Arial", size=10.5, bold=False, italic=False, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_para(doc, text="", size=10.5, bold=False, italic=False, color=None,
             align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=6, keep=False):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.keep_with_next = keep
    font(p.add_run(text), size=size, bold=bold, italic=italic, color=color)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.add_run(text)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.12
    font(p.add_run(text), size=10.5)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True, color=WHITE, size=9.2,
                      align=WD_ALIGN_PARAGRAPH.CENTER)
        shade(table.rows[0].cells[i], NAVY)
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, size=9.1,
                          align=WD_ALIGN_PARAGRAPH.CENTER if i in (0, len(row)-1) else WD_ALIGN_PARAGRAPH.LEFT)
            if ridx % 2:
                shade(cells[i], PALE)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
                tc_pr = row.cells[i]._tc.get_or_add_tcPr()
                tc_w = tc_pr.find(qn("w:tcW"))
                if tc_w is not None:
                    tc_w.set(qn("w:w"), str(int(width * 1440)))
                    tc_w.set(qn("w:type"), "dxa")
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_code(doc, text, caption):
    add_para(doc, caption, size=9.5, bold=True, color=GRAY, align=WD_ALIGN_PARAGRAPH.LEFT, after=3)
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    shade(cell, "F2F4F7")
    margins(cell, 110, 160, 110, 160)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(text)
    font(r, name="Consolas", size=8.5, color="263238")
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def page_break(doc):
    doc.add_page_break()


doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.8)
sec.bottom_margin = Inches(0.75)
sec.left_margin = Inches(0.9)
sec.right_margin = Inches(0.9)
sec.header_distance = Inches(0.3)
sec.footer_distance = Inches(0.35)

# Standard business brief tokens, adapted to an academic weekly report.
styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.15
for level, size, before, after in ((1, 16, 14, 7), (2, 13, 10, 5), (3, 11.5, 8, 4)):
    s = styles[f"Heading {level}"]
    s.font.name = "Arial"
    s._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    s.font.size = Pt(size)
    s.font.bold = True
    s.font.color.rgb = RGBColor.from_string(BLUE if level < 3 else NAVY)
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)

# Running header/footer.
hp = sec.header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
font(hp.add_run("FER202  |  ATELIER  |  BÁO CÁO TUẦN 10"), size=8.5, bold=True, color=GRAY)
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(fp.add_run("Nhóm FER202 — rbl-fer202-theotuan-36-38-75_gr"), size=8, color=GRAY)

# Cover: editorial cover.
add_para(doc, "TRƯỜNG ĐẠI HỌC FPT", 12, True, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
add_para(doc, "KHOA CÔNG NGHỆ THÔNG TIN", 10.5, True, color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER, after=70)
add_para(doc, "BÁO CÁO TIẾN ĐỘ DỰ ÁN", 10, True, color=GOLD, align=WD_ALIGN_PARAGRAPH.CENTER, after=12)
add_para(doc, "TUẦN 10", 30, True, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
add_para(doc, "Nâng cao, tối ưu và triển khai", 15, False, color=BLUE, align=WD_ALIGN_PARAGRAPH.CENTER, after=22)
add_para(doc, "ATELIER — LUXURY SUPERCAR SHOWROOM", 16, True, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
add_para(doc, "Ứng dụng ReactJS trưng bày và quản lý siêu xe hạng sang", 10.5, italic=True, color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER, after=55)
add_table(doc, ["THÔNG TIN", "NỘI DUNG"], [
    ["Môn học", "Phát triển ứng dụng Web với ReactJS (FER202)"],
    ["Nhóm", "5 thành viên — mã repository 36-38-75"],
    ["Tuần báo cáo", "Tuần 10 — Redux, Lazy Loading, Tailwind và Deploy"],
    ["Ngày hoàn thiện", "21/07/2026"],
    ["Repository", "github.com/fudn-traltb-su26/rbl-fer202-theotuan-36-38-75_gr"],
], [1.5, 4.8])

page_break(doc)
add_heading(doc, "I. THÔNG TIN VÀ MỤC TIÊU TUẦN 10", 1)
add_heading(doc, "1.1. Tổng quan dự án", 2)
add_para(doc, "ATELIER là nền tảng trưng bày và khám phá siêu xe hạng sang được xây dựng bằng React 18 và Vite. Hệ thống cung cấp danh sách xe, chi tiết thông số, so sánh xe, tìm kiếm, danh sách yêu thích, liên hệ đặt lịch và trang quản trị. Giao diện sử dụng phong cách Dark Mode kết hợp điểm nhấn màu vàng champagne.")
add_heading(doc, "1.2. Mục tiêu", 2)
for item in [
    "Chuyển phần trạng thái dùng chung sang Redux Toolkit để dữ liệu nhất quán và dễ mở rộng.",
    "Áp dụng React.lazy() và Suspense cho toàn bộ 14 trang nhằm tạo các JavaScript chunk riêng.",
    "Tích hợp Tailwind CSS vào pipeline PostCSS và sử dụng utility class ở phạm vi phù hợp.",
    "Tạo production build bằng Vite và cấu hình Vercel hỗ trợ React Router.",
    "Kiểm tra lại trải nghiệm tải trang, chức năng yêu thích, tìm kiếm và các route quan trọng.",
]: add_bullet(doc, item)
add_heading(doc, "1.3. Phạm vi và tiêu chí hoàn thành", 2)
add_table(doc, ["Hạng mục", "Tiêu chí", "Kết quả"], [
    ["Redux Toolkit", "Store, slice, action, selector và Provider hoạt động", "Đạt"],
    ["Lazy Loading", "14 page component được dynamic import và có fallback", "Đạt"],
    ["Tailwind CSS", "Có cấu hình content, PostCSS và directives toàn cục", "Đạt"],
    ["Production build", "npm run build không có lỗi", "Đạt"],
    ["Vercel", "Có rewrite SPA về index.html", "Đạt"],
], [1.55, 3.75, 1.0])

add_heading(doc, "II. PHÂN CÔNG CÔNG VIỆC", 1)
add_table(doc, ["STT", "Thành viên", "Vai trò Tuần 10", "% HT"], [
    ["1", "Văn Anh", "Trưởng nhóm; kiểm tra routing, tích hợp và review", "100%"],
    ["2", "Sang", "Thiết lập Redux Toolkit, store và carSlice", "100%"],
    ["3", "Tuân", "Lazy loading, Tailwind CSS, build và Vercel", "100%"],
    ["4", "Sơn", "Kiểm tra dữ liệu/API và các trang hiển thị", "100%"],
    ["5", "Quý", "Kiểm thử state yêu thích/tìm kiếm và tài liệu", "100%"],
], [0.45, 1.1, 4.0, 0.65])
add_para(doc, "Người thực hiện chính theo TODO.md: Sang và Tuân. Các thành viên còn lại tham gia review, kiểm thử tích hợp và xác nhận đầu ra.", italic=True, color=GRAY)

page_break(doc)
add_heading(doc, "III. NỘI DUNG ĐÃ THỰC HIỆN", 1)
add_heading(doc, "3.1. Quản lý trạng thái bằng Redux Toolkit", 2)
add_para(doc, "Nhóm tạo Redux Store tại src/store/store.js và đăng ký carReducer dưới khóa car. File carSlice.js quản lý hai trạng thái dùng chung là favorites và searchQuery. Danh sách yêu thích được đồng bộ với localStorage để vẫn tồn tại sau khi tải lại trang. Các action và selector được export để những lớp giao diện có thể dispatch hoặc đọc state một cách thống nhất.")
add_code(doc, "export const store = configureStore({\n  reducer: { car: carReducer },\n});", "Minh chứng 1 — Cấu hình Redux Store")
add_code(doc, "const carSlice = createSlice({\n  name: 'car',\n  initialState: { favorites: loadFavorites(), searchQuery: '' },\n  reducers: { toggleFavorite, setFavorites, setSearchQuery },\n});", "Minh chứng 2 — Cấu trúc carSlice")
add_para(doc, "Trong main.jsx, ứng dụng được bọc bởi <Provider store={store}>. CarContext vẫn được duy trì như một lớp tương thích để các component đang dùng useCar() không phải thay đổi đồng loạt; dữ liệu bên trong context đã lấy từ useSelector và cập nhật bằng useDispatch.")

add_heading(doc, "3.2. Lazy Loading và Suspense", 2)
add_para(doc, "Các import tĩnh của 14 trang được thay bằng React.lazy(). Routes được đặt bên trong Suspense với giao diện fallback có spinner và thông báo tải. Khi người dùng truy cập một route, trình duyệt mới tải chunk JavaScript tương ứng, từ đó giảm lượng mã cần tải ngay ở lần mở đầu tiên.")
add_code(doc, "const HomePage = lazy(() => import('./pages/HomePage'))\nconst CarDetailPage = lazy(() => import('./pages/CarDetailPage'))\nconst AdminDashboard = lazy(() => import('./pages/AdminDashboard'))\n\n<Suspense fallback={<LoadingView />}>\n  <Routes>...</Routes>\n</Suspense>", "Minh chứng 3 — Dynamic import và Suspense trong App.jsx")
add_table(doc, ["Nhóm route", "Các trang đã lazy load"], [
    ["Khám phá", "Home, Cars, Car Detail, Compare, Brands, News, Gallery"],
    ["Người dùng", "Login, Register, Favorites, About, Search, Contact"],
    ["Quản trị", "Admin Dashboard (được bảo vệ bởi ProtectedRoute)"],
], [1.45, 4.85])

page_break(doc)
add_heading(doc, "3.3. Tích hợp Tailwind CSS", 2)
add_para(doc, "Tailwind CSS và PostCSS đã được thêm vào package.json. File tailwind.config.js khai báo vùng quét index.html và toàn bộ file JavaScript/JSX trong src. Ba directive Tailwind được đặt ở đầu src/index.css. Nhóm áp dụng utility class có chọn lọc và tiếp tục giữ Bootstrap cùng custom CSS để tránh làm thay đổi toàn bộ thiết kế Dark Mode hiện có.")
add_code(doc, "content: [\n  './index.html',\n  './src/**/*.{js,ts,jsx,tsx}',\n]\n\n@tailwind base;\n@tailwind components;\n@tailwind utilities;", "Minh chứng 4 — Tailwind content scan và CSS directives")

add_heading(doc, "3.4. Production build và Vercel", 2)
add_para(doc, "Nhóm chạy npm run build bằng Vite. Kết quả thực tế ngày 21/07/2026: 458 module được biến đổi, quá trình build hoàn tất trong 4,34 giây và không có lỗi. Thư mục dist chứa nhiều tệp JavaScript riêng cho từng trang, cho thấy cơ chế code splitting đã hoạt động.")
add_table(doc, ["Chỉ số kiểm tra", "Kết quả"], [
    ["Câu lệnh", "npm run build"],
    ["Vite", "v5.4.21"],
    ["Modules transformed", "458"],
    ["Thời gian build", "4,34 giây"],
    ["Trạng thái", "Thành công — không có error"],
    ["Dấu hiệu code splitting", "Có chunk riêng cho HomePage, CarDetailPage, AdminDashboard..."],
], [2.0, 4.3])
add_code(doc, '{\n  "rewrites": [\n    { "source": "/(.*)", "destination": "/index.html" }\n  ]\n}', "Minh chứng 5 — vercel.json hỗ trợ SPA routing")

add_heading(doc, "3.5. Quy trình thực hiện", 2)
for item in [
    "Khảo sát state dùng chung trong CarContext và xác định favorites/searchQuery cần chuyển sang Redux.",
    "Cài đặt dependencies, tạo store và slice, sau đó bọc ứng dụng bằng Redux Provider.",
    "Chuyển các page import sang dynamic import; bổ sung Suspense fallback.",
    "Cấu hình Tailwind/PostCSS và Vercel rewrite cho React Router.",
    "Chạy production build, kiểm tra các chunk đầu ra và rà soát checklist Tuần 10.",
]: add_bullet(doc, item)

page_break(doc)
add_heading(doc, "IV. KẾT QUẢ VÀ MINH CHỨNG GIAO DIỆN", 1)
add_para(doc, "Sau khi tích hợp, dự án vẫn giữ được ngôn ngữ thiết kế ATELIER và bổ sung nền tảng kỹ thuật phục vụ mở rộng. Hai hình dưới đây thể hiện các giao diện chính được sử dụng trong quá trình kiểm tra và hoàn thiện.")
for img_path, caption in [
    (ROOT / "docs" / "mockup" / "HomePage" / "screen.png", "Hình 1. Giao diện HomePage của ATELIER — trang được tải theo cơ chế lazy loading."),
    (ROOT / "docs" / "mockup" / "AdminDashboard" / "screen.png", "Hình 2. Giao diện Admin Dashboard — route quản trị được lazy load và bảo vệ."),
]:
    if img_path.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run()
        r.add_picture(str(img_path), width=Inches(5.8))
        add_para(doc, caption, size=9, italic=True, color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER, after=10)

add_heading(doc, "4.1. Kết quả kỹ thuật", 2)
for item in [
    "State yêu thích và tìm kiếm được quản lý tập trung bằng Redux Toolkit.",
    "Danh sách yêu thích được lưu trong localStorage và phục hồi khi khởi động ứng dụng.",
    "14 trang được tách thành các chunk tải theo nhu cầu.",
    "Production build hoàn thành thành công với 458 module.",
    "Vercel rewrite ngăn lỗi 404 khi tải lại route con.",
]: add_bullet(doc, item)

page_break(doc)
add_heading(doc, "V. CHECKLIST ĐẦU RA", 1)
add_table(doc, ["#", "Đầu ra", "HT?", "Vị trí minh chứng"], [
    ["1", "Redux Store và carSlice", "Có", "src/store/store.js; src/store/slices/carSlice.js"],
    ["2", "Redux Provider", "Có", "src/main.jsx"],
    ["3", "React.lazy cho 14 pages", "Có", "src/App.jsx"],
    ["4", "Suspense fallback", "Có", "src/App.jsx"],
    ["5", "Tailwind và PostCSS", "Có", "tailwind.config.js; postcss.config.js; src/index.css"],
    ["6", "Production build", "Có", "dist/ — build ngày 21/07/2026"],
    ["7", "Vercel SPA rewrite", "Có", "vercel.json"],
    ["8", "TODO Tuần 10", "Có", "TODO.md — các mục đã đánh dấu [x]"],
], [0.35, 2.55, 0.55, 2.85])
add_para(doc, "Tổng số đầu ra hoàn thành: 8/8.", bold=True, color=NAVY, align=WD_ALIGN_PARAGRAPH.LEFT)

add_heading(doc, "VI. KHÓ KHĂN, HẠN CHẾ VÀ HƯỚNG PHÁT TRIỂN", 1)
add_heading(doc, "6.1. Khó khăn và cách xử lý", 2)
add_table(doc, ["Vấn đề", "Nguyên nhân", "Cách xử lý", "Trạng thái"], [
    ["Nguy cơ phải sửa nhiều component khi đổi state", "Các component đang dùng useCar()", "Giữ CarContext làm adapter nối sang Redux", "Đã xử lý"],
    ["Trang tải lần đầu có bundle lớn", "Import tĩnh toàn bộ pages", "React.lazy và Suspense", "Đã xử lý"],
    ["Reload route con trên hosting", "Server không biết route phía client", "Rewrite mọi route về index.html", "Đã xử lý"],
    ["Nhiều hệ thống CSS cùng tồn tại", "Bootstrap, custom CSS và Tailwind", "Chỉ migration có chọn lọc, tránh xung đột", "Theo dõi"],
], [1.55, 1.45, 2.45, 0.85])
add_heading(doc, "6.2. Hạn chế", 2)
for item in [
    "CarContext chưa được loại bỏ hoàn toàn; hiện là lớp adapter để bảo đảm tương thích.",
    "Tailwind mới được tích hợp ở mức nền tảng, phần lớn giao diện vẫn dùng Bootstrap và custom CSS.",
    "Báo cáo chưa ghi URL production vì repository không lưu đường dẫn Vercel công khai.",
    "Cần bổ sung kiểm thử tự động cho reducer và hành vi route khi dự án tiếp tục phát triển.",
]: add_bullet(doc, item)
add_heading(doc, "6.3. Hướng phát triển", 2)
for item in [
    "Chuyển dần component từ useCar() sang useSelector/useDispatch trực tiếp rồi loại bỏ adapter.",
    "Chuẩn hóa styling để giảm chồng chéo giữa Bootstrap, Tailwind và custom CSS.",
    "Bổ sung unit test cho carSlice và integration test cho Favorites/Search.",
    "Theo dõi kích thước bundle và tối ưu thêm ảnh/video của showroom.",
]: add_bullet(doc, item)

add_heading(doc, "VII. KẾT LUẬN", 1)
add_para(doc, "Tuần 10 đã hoàn thành các mục tiêu trọng tâm về quản lý trạng thái, tối ưu tải và triển khai. Redux Toolkit giúp state dùng chung có cấu trúc rõ ràng; lazy loading tạo các chunk theo trang; Tailwind được tích hợp vào pipeline; production build chạy thành công; và Vercel rewrite bảo đảm routing phía client. Kết quả này đưa dự án ATELIER từ giai đoạn phát triển tính năng sang trạng thái sẵn sàng trình diễn và triển khai production.")
add_para(doc, "Commit minh chứng chính: d6b4d5b — hoàn thiện Issue/TODO Tuần 10; thay đổi 25 file với 1.989 dòng bổ sung và 168 dòng loại bỏ.", bold=True, color=NAVY, align=WD_ALIGN_PARAGRAPH.LEFT)

doc.core_properties.title = "Báo cáo Tuần 10 — ATELIER"
doc.core_properties.subject = "FER202 — Redux Toolkit, Lazy Loading, Tailwind CSS và Vercel"
doc.core_properties.author = "Nhóm FER202 — ATELIER"
doc.core_properties.keywords = "FER202, ReactJS, Redux Toolkit, Lazy Loading, Tailwind, Vercel"
doc.save(OUT)
print(OUT)
